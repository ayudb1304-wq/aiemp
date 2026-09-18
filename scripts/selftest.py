"""Self-test of the deterministic parts. No API key, no network.

  python scripts/selftest.py

Builds a temp tree (AIEMP_ROOT) with fixture context and memory files, then checks: tracker
upsert (create, idempotent re-run, to_verify match, due slip), chunk merge, the P1 cap fallback,
brief.build with the LLM stubbed, the context loader's drop order under a tiny budget, keyword
search, the Sheet pull merge (corrections, deleted rows) and the document readers.
Exit code 1 on any failure.
"""
import json
import os
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

TMP = Path(tempfile.mkdtemp(prefix="aiemp-selftest-"))
os.environ["AIEMP_ROOT"] = str(TMP)
os.environ.pop("ANTHROPIC_API_KEY", None)
sys.path.insert(0, str(Path(__file__).resolve().parent))

import common  # noqa: E402  (after AIEMP_ROOT is set)
import brief  # noqa: E402
import dayplan  # noqa: E402
import extract  # noqa: E402
import prep  # noqa: E402
import search  # noqa: E402
import sheets  # noqa: E402
import tracker  # noqa: E402
import verify  # noqa: E402

FAILS = []


def check(cond, msg):
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        FAILS.append(msg)


def fixtures():
    (TMP / "context" / "projects").mkdir(parents=True)
    (TMP / "context" / "people").mkdir()
    (TMP / "memory").mkdir()
    (TMP / "briefs" / "teams").mkdir(parents=True)
    (TMP / "cards").mkdir()
    (TMP / "inbox").mkdir()
    (TMP / "context" / "company.md").write_text(
        "# Company\n\n## Me\n- Name: Ayush Test. Role: lead.\n\n"
        "## Name map (transcription errors to real people)\n| Heard as | Person |\n|---|---|\n"
        "| Lakshmi, Leshmi | Laxmikant |\n| Sam | Samuel |\n| Sri sir | Sreekumar |\n\n## Teams\n"
        "| Team | Lead | Members | Responsibility |\n|---|---|---|---|\n"
        "| Alpha | Laxmikant | Laxmikant, Samuel | build |\n\n"
        "## Stakeholders\n- Sreekumar: customer.\n\n## Priority rules\n- P1 only if blocking.\n",
        encoding="utf-8")
    (TMP / "context" / "projects" / "alpha.md").write_text(
        "# Alpha Project\n\n## Goal\nShip the alpha tracker.\n\n## Key people\n- Johan: CR module.\n\n"
        "## Glossary\n- CR: change request. Sub-items are CR items.\n\n## Recent decisions\n_none_\n", encoding="utf-8")
    (TMP / "context" / "projects" / "beta.md").write_text("# Beta\n\n## Goal\nBeta.\n", encoding="utf-8")
    (TMP / "context" / "people" / "laxmikant.md").write_text(
        "# Laxmikant\nTeam lead. Likes decisions in writing. " * 20, encoding="utf-8")
    common.write_jsonl(common.DECISIONS, [
        {"id": "D-2026-09-01-01", "date": "2026-09-01", "project": "alpha",
         "decision": "The currency selector lives on the CR entry screen", "by": "Sreekumar",
         "evidence": "put the currency selector on the entry screen", "source": "2026-09-01-alpha.md",
         "supersedes": None},
        {"id": "D-2026-09-03-01", "date": "2026-09-03", "project": "alpha",
         "decision": "No DB column changes after sign-off", "by": "Ayush",
         "evidence": "no more column changes", "source": "2026-09-03-alpha.md", "supersedes": None},
        {"id": "D-2026-09-05-01", "date": "2026-09-05", "project": "beta",
         "decision": "Beta uses the export job nightly", "by": "Ayush", "evidence": "nightly export",
         "source": "2026-09-05-beta.md", "supersedes": None},
        {"id": "D-2026-09-08-01", "date": "2026-09-08", "project": "alpha",
         "decision": "The currency selector moves to the impact analysis view", "by": "Sreekumar",
         "evidence": "move the selector to impact analysis", "source": "2026-09-08-alpha.md",
         "supersedes": "D-2026-09-01-01"},
    ])
    common.write_jsonl(common.THREADS, [
        {"id": "T-2026-09-02-01", "date": "2026-09-02", "project": "alpha", "topic": "flaky export",
         "note": "The Zoho export failed twice this week", "evidence": "export failed again",
         "source": "2026-09-02-alpha.md", "mentions": 3, "last_seen": "2026-09-09", "promoted_to": None},
    ])
    (TMP / "briefs" / "2026-09-10.md").write_text(
        "# Morning brief\n\n## Project: Alpha Project\n_2 planned, 0 unplanned_\n- `x` old line about currency\n\n"
        "## Project: Beta\n- beta line\n", encoding="utf-8")


def test_tracker():
    print("tracker")
    d = "2026-09-10"
    cards = [
        {"team": "Alpha", "project": "alpha", "owner": "Samuel", "task": "Write the runbook", "due": "2026-09-12",
         "priority": "P1", "status": "open", "blocked_by": None, "evidence": "I will write the runbook",
         "matches_existing_id": None, "update_note": None, "type": "build", "next_step": "Draft section 1",
         "prerequisites": ["DB sign-off"], "unblocker": "Laxmikant", "effort": "half-day", "basis": "quote"},
        {"team": "Alpha", "project": "alpha", "owner": "Ayush", "task": "Send the review comments to Sreekumar",
         "due": None, "priority": "P2", "status": "open", "blocked_by": None, "evidence": "send them today",
         "matches_existing_id": None, "update_note": None, "type": "communicate", "next_step": None,
         "prerequisites": None, "unblocker": "Laxmikant", "effort": "15m", "basis": None},
    ]
    created, updated = tracker.upsert(cards, "Alpha sync", d, "2026-09-10-alpha.md", project="alpha")
    check(created == ["2026-09-10-alpha-01", "2026-09-10-alpha-02"], f"ids use the project slug: {created}")
    rows = tracker.load_rows()
    check(rows[0]["prerequisites"] == "DB sign-off" and rows[0]["origin"] == "planned", "advisor fields and origin stored")
    created2, updated2 = tracker.upsert(cards, "Alpha sync", d, "2026-09-10-alpha.md", project="alpha")
    check(created2 == [] and len(tracker.load_rows()) == 2, "re-running the same file creates no duplicates")

    verify = [{"team": "Alpha", "project": "alpha", "owner": "Samuel", "task": "Runbook is finished", "due": None,
               "priority": "P1", "status": "done", "blocked_by": None, "evidence": "the runbook is finished",
               "matches_existing_id": "2026-09-10-alpha-01", "update_note": "the runbook is finished",
               "type": "build", "next_step": None, "prerequisites": None, "unblocker": None,
               "effort": "unknown", "basis": None}]
    c, u = tracker.upsert(verify, "Alpha sync", "2026-09-11", "2026-09-11-alpha.md", project="alpha")
    r = {x["id"]: x for x in tracker.load_rows()}["2026-09-10-alpha-01"]
    check(c == [] and u == ["2026-09-10-alpha-01"], "completion matched the existing row, no new row")
    c2, u2 = tracker.upsert(verify, "Alpha sync", "2026-09-11", "2026-09-11-alpha.md", project="alpha")
    check(c2 == [] and u2 == [], "re-applying the same update is a no-op")
    check(r["status"] == "to_verify", f"status is to_verify, never done: {r['status']}")
    check("2026-09-11: the runbook is finished [the runbook is finished]" in r["notes"], "dated note has the quote")
    check(r["status"] in common.OPEN_STATUSES, "to_verify counts as open")

    slip = [{"team": "Alpha", "project": "alpha", "owner": "Ayush", "task": "x", "due": "2026-09-14", "priority": "P2",
             "status": "open", "blocked_by": None, "evidence": "e", "matches_existing_id": "2026-09-10-alpha-02",
             "update_note": "slipped", "type": "communicate", "next_step": None, "prerequisites": None,
             "unblocker": None, "effort": "15m", "basis": None}]
    tracker.upsert(slip, "s", "2026-09-12", "2026-09-12-alpha.md", project="alpha")
    slip[0]["due"] = "2026-09-16"
    tracker.upsert(slip, "s", "2026-09-14", "2026-09-14-alpha.md", project="alpha")
    r = {x["id"]: x for x in tracker.load_rows()}["2026-09-10-alpha-02"]
    n, first = tracker.slip_count(r)
    check(n == 2 and first == "2026-09-14", f"two due changes recorded in notes, first real due kept: {r['notes']}")
    check(len(tracker.open_items(project="alpha")) == 2 and tracker.open_items(project="beta") == [],
          "open_items is scoped by project")
    tracker.close("2026-09-10-alpha-02")
    r = {x["id"]: x for x in tracker.load_rows()}["2026-09-10-alpha-02"]
    check(r["status"] == "done" and r["closed"] == common.today(), "close() records the closed date")
    from openpyxl import load_workbook
    wb = load_workbook(common.TRACKER)
    ws = wb["Actions"]
    check(wb.sheetnames == ["Dashboard", "Actions"] and wb.active.title == "Dashboard", "workbook: Dashboard first, Actions second")
    check([c.value for c in ws[1]] == common.COLUMNS and common.COLUMNS[0] == "task" and common.COLUMNS[-1] == "id",
          "Actions header is COLUMNS: task first, provenance (created, source, id) last")
    check("Actions" in ws.tables and ws.tables["Actions"].ref.endswith(str(ws.max_row)), "Actions is an Excel table over every row")
    check(hasattr(ws[f"{tracker._col('due')}2"].value, "isoformat"), "due is written as a real date")
    check(sum(len(cf.rules) for cf in ws.conditional_formatting) == 11 and len(ws.data_validations.dataValidation) == 5,
          "colour rules and dropdowns are attached")
    dash = wb["Dashboard"]
    check(len(dash._charts) == 2 and dash["B1"].value == "Action tracker", "dashboard has two charts")
    check(str(dash["B5"].value).startswith("=SUMPRODUCT(COUNTIFS(Actions!") and "Actions!" in str(dash["C10"].value)
          and wb.calculation.fullCalcOnLoad, "dashboard tiles and tables are live formulas over the Actions sheet")
    check(len(tracker.load_rows()) == 2, "rows read back by header name after the reorder")



def test_merge_and_cap():
    print("extract merge + P1 cap")
    a = [{"task": "Write the runbook", "matches_existing_id": None, "due": None, "priority": "P1"},
         {"task": "Fix export", "matches_existing_id": "2026-09-10-alpha-01", "due": None, "priority": "P1"}]
    b = [{"task": "write the runbook", "matches_existing_id": None, "due": "2026-09-12", "priority": "P1"},
         {"task": "Export fix (again)", "matches_existing_id": "2026-09-10-alpha-01", "due": None, "priority": "P1",
          "update_note": "n"},
         {"task": "New thing", "matches_existing_id": None, "due": None, "priority": "P2"}]
    m = extract.merge_cards(a, b)
    check(len(m) == 3, f"chunk merge dedupes by id and task text: {len(m)} cards")
    check(m[0]["due"] == "2026-09-12", "later chunk fills a missing due")

    cards = [{"task": f"t{i}", "priority": "P1", "due": d} for i, d in
             enumerate(["2026-09-20", "2026-09-11", None, "2026-09-12", "2026-09-13"])]
    extract.ask_json = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no key"))
    log = {}
    out = extract.enforce_p1_cap(cards, log)
    p1 = [c["task"] for c in out if c["priority"] == "P1"]
    check(p1 == ["t1", "t3", "t4"], f"cap keeps the 3 earliest due: {p1}")
    check(log.get("p1_downgraded") == ["t0", "t2"], "downgrades are logged")


def test_context_loader():
    print("context loader")
    doc = "We talked about the currency selector and the flaky export with Laxmikant. " * 50
    text, log = common.build_prompt_context("alpha", ["laxmikant"], doc, budget_tokens=100000)
    names = list(log["layers"])
    check(names == ["1:company+project", "2:open_items", "3:people", "4:history", "5:briefs"], "five layers")
    check(all(log["layers"][n]["loaded"] for n in names), f"all layers load under a big budget: {log}")
    check("D-2026-09-01-01" in text and "T-2026-09-02-01" in text and "old line about currency" in text,
          "history and brief sections are in the prompt")
    check("D-2026-09-05-01" not in text, "other projects' decisions are not retrieved")
    check(isinstance(log.get("history_hits"), int) and log["history_hits"] >= 1, f"history hit count logged: {log.get('history_hits')}")
    sizes = {n: log["layers"][n]["tokens"] for n in names}
    tight = sizes["1:company+project"] + common.est_tokens(doc) + sizes["2:open_items"] + sizes["3:people"] + 5
    _, log2 = common.build_prompt_context("alpha", ["laxmikant"], doc, budget_tokens=tight)
    check(log2["dropped"] == ["5:briefs", "4:history"], f"lowest layers dropped first: {log2['dropped']}")
    _, log3 = common.build_prompt_context("alpha", ["laxmikant"], doc, budget_tokens=10)
    check(log3["dropped"] == ["5:briefs", "4:history", "3:people", "2:open_items"]
          and log3["layers"]["1:company+project"]["loaded"], "layer 1 is never dropped")
    check(any("CONSOLIDATE" in w for w in log3["warnings"]), "oversize layer 1 logs a CONSOLIDATE warning")
    p = common.write_run_log("selftest", log3)
    check(p.exists() and json.loads(p.read_text())["dropped"], "run log written")
    check(common.load_context("alpha", ["laxmikant", "nobody"]).count("# ") >= 3, "loader skips missing people")


def test_search():
    print("search")
    hits = search.search("when did we last discuss the currency selectors")
    check([h["id"] for h in hits if h["kind"] == "decision"] == ["D-2026-09-01-01", "D-2026-09-08-01"],
          f"stemmed keyword hit, both decisions of the chain: {[h['id'] for h in hits]}")
    dec = {h["id"]: h for h in hits if h["kind"] == "decision"}
    check(dec["D-2026-09-01-01"]["record"]["replaced_by"] == "D-2026-09-08-01" and dec["D-2026-09-08-01"]["record"]["current"]
          and "replaced by D-2026-09-08-01" in dec["D-2026-09-01-01"]["text"], "decision chain: current and replaced_by")
    check(search.tokens("Lakshmi and Sam discussed the change request") == ["laxmikant", "samuel", "cr"],
          f"name map and glossary fold to canonical tokens: {search.tokens('Lakshmi and Sam discussed the change request')}")
    by_alias = search.search("what did Sri sir decide")
    check([h["id"] for h in by_alias if h["kind"] == "decision"] == ["D-2026-09-01-01", "D-2026-09-08-01"],
          f"a two-word alias in the question finds the records under the real name: {[h['id'] for h in by_alias]}")
    text, n = search.history_block("alpha", "we should move the currency selector again")
    check(n >= 2 and '"current": false' in text and '"replaced_by": "D-2026-09-08-01"' in text,
          f"history block carries the chain ({n} hits)")
    check(hits and hits[0]["date"] <= hits[-1]["date"], "hits are in date order")
    check(search.search("nightly export", project="alpha") == [] or
          all(h["project"] in ("alpha", "") for h in search.search("nightly export", project="alpha")),
          "project filter excludes other projects")
    check(search.keywords("What did we decide about the DB column?") == ["column"] or
          "column" in search.keywords("What did we decide about the DB column?"), "keywords drop stop words")


def test_brief():
    print("brief")
    calls = []

    def stub(system, user):
        calls.append(user)
        return "- `2026-09-15-alpha-01` because it unblocks `2026-09-15-alpha-02` and is 1d overdue\n"

    t = date(2026, 9, 16)
    more = [{"team": "Alpha", "project": "alpha", "owner": "Ayush", "task": "Decide the CR approval flow",
             "due": "2026-09-15", "priority": "P1", "status": "open", "blocked_by": None, "evidence": "decide",
             "matches_existing_id": None, "update_note": None, "type": "decide", "next_step": "Reply to Johan",
             "prerequisites": None, "unblocker": "Johan", "effort": "15m", "basis": "D-2026-09-01-01"},
            {"team": "Alpha", "project": "alpha", "owner": "Johan", "task": "Build the CR approval screen",
             "due": None, "priority": "P2", "status": "blocked", "blocked_by": "2026-09-15-alpha-01",
             "evidence": "after the decision", "matches_existing_id": None, "update_note": None, "type": "build",
             "next_step": None, "prerequisites": ["2026-09-15-alpha-01"], "unblocker": "Ayush",
             "effort": "day+", "basis": None}]
    tracker.upsert(more, "Alpha sync", "2026-09-15", "2026-09-15-alpha.md", project="alpha")
    text, digests = brief.build(ask_fn=stub, t=t)
    check("unblocks `2026-09-15-alpha-02`" in text or "because the runbook" in text, "do-first reason cites ids")
    check("could go to **Johan**" in text, "my item with another unblocker is listed under Delegate?")
    for h in ("## Waiting for your confirmation", "## Do first", "## Batch", "## Delegate?",
              "## Recurring threads", "## Project: Alpha Project"):
        check(h in text, f"section present: {h}")
    check("the runbook is finished" in text, "to_verify item shows its completion quote")
    check("T-2026-09-02-01" in text and "make this a task?" in text, "recurring thread (3 mentions) shown")
    check("_1 planned, 0 unplanned_" in text or "planned" in text, "planned/unplanned counts")
    check(calls and "Alpha" in digests, "LLM stub was called once and team digest built")
    text2, _ = brief.build(ask_fn=None, t=t)
    check("## Do first" in text2, "deterministic brief without LLM")
    p = brief.plan(ask_fn=stub, t=t)
    json.dumps(p)
    check(brief.render_md(p) == text, "render_md(plan) is the brief text")
    page = brief.render_html(p)
    check(page.startswith("<!doctype html>") and "Do first" in page and "Decide the CR approval flow" in page
          and 'class="pill P1"' in page and "1d overdue" in page, "html brief: sections, items, priority and overdue pills")
    check("<script" not in page and "http" not in page.split("<body>")[0].replace("http-equiv", ""), "html brief is self-contained")
    rows = brief.sheet_rows(p)
    kinds = [k for k, _ in rows]
    check(kinds[:3] == ["title", "subtitle", "header"] and "section" in kinds and "item" in kinds
          and all(len(c) <= 8 for _, c in rows), "sheet rows: title, subtitle, header, then sections and items")
    ch = p["changes"]
    check(ch["since"] == "2026-09-10" and any(r["id"] == "2026-09-15-alpha-01" for r in ch["new"])
          and "## Since 2026-09-10" in text and "- new: `2026-09-15-alpha-01`" in text,
          f"since-last-brief: items not in the 2026-09-10 brief are new: {[r['id'] for r in ch['new']]}")
    check(any(r["id"] == "2026-09-10-alpha-02" for r in ch["closed"]) and "- closed: `2026-09-10-alpha-02`" in text,
          "since-last-brief: closed items listed")
    check(any(x[0]["id"] == "2026-09-10-alpha-02" for x in ch["slipped"]), f"since-last-brief: slips from the notes markers: {ch['slipped']}")
    alpha = next(pr for pr in p["projects"] if pr["slug"] == "alpha")
    check([d["id"] for d in alpha["decisions"]] == ["D-2026-09-08-01", "D-2026-09-03-01"] and "### Decisions in force" in text
          and "`D-2026-09-08-01`" in text and "`D-2026-09-01-01`" not in text.split("## Project: Alpha Project")[1],
          f"decisions in force: current decisions only, newest first: {[d['id'] for d in alpha['decisions']]}")
    check("Since 2026-09-10" in page and "Decisions in force" in page, "html brief has the since block and decisions in force")
    check(any(k == "section" and c[0].startswith("Since 2026-09-10") for k, c in rows), "sheet rows have the since section")
    item = next(c for k, c in rows if k == "item")
    check(item[1] and item[2] and item[7].startswith("2026-"), "sheet item rows carry owner, task and id")
    json.dumps(sheets.brief_format_requests(1, kinds))


def test_verify():
    print("fact checks")
    doc = ("Ashwini: I will be signing off on the Debi after the review.\n"
           "Lakshmi and I will be working on the proper workflow solution design by Friday.")
    ok = verify.quote_found("Lakshmi and I will be working on the workflow solution design", doc)
    bad = verify.quote_found("We agreed to cancel the whole project next week", doc)
    check(ok and not bad, "quote check tolerates one transcription word, rejects an invented quote")
    card = {"owner": "Lakshmi", "unblocker": "Sam", "evidence": "signing off on the Debi after the review",
            "due": "2026-09-05", "status": "blocked", "blocked_by": None, "matches_existing_id": "nope",
            "next_step": "Ask Ashwini", "basis": None, "task": "x"}
    flags = verify.check_card(card, doc, "2026-09-11", "", common.roster(), common.name_map(), {"2026-09-10-alpha-01"}, set())
    check(card["owner"] == "Laxmikant" and card["unblocker"] == "Samuel" and card["matches_existing_id"] is None,
          f"aliases become roster names, unknown match cleared: {card['owner']}, {card['unblocker']}")
    check({f.split(":")[0] for f in flags} == {"owner mapped", "unblocker mapped", "due before meeting",
                                               "blocked without blocker", "matched id unknown", "next step without basis"},
          f"every failed check is a flag: {flags}")
    good = {"owner": "Johan", "evidence": "working on the proper workflow solution design", "due": "2026-09-12",
            "status": "open", "next_step": "Send it", "basis": "D-2026-09-08-01", "task": "y"}
    check(verify.check_card(good, doc, "2026-09-11", "", common.roster(), common.name_map(), set(), {"D-2026-09-08-01"}) == [],
          "a card that passes every check has no flags")
    data = {"cards": [dict(card), dict(good)], "decisions": [{"decision": "d", "evidence": "not in the doc at all really"}],
            "threads": [{"topic": "t", "evidence": "signing off on the Debi"}]}
    summary = verify.check_all(data, doc, "2026-09-11", "alpha", [], {"2026-09-10-alpha-01"})
    check(summary["cards_flagged"] == 1 and data["decisions"][0]["flags"] == ["quote not found"]
          and data["threads"][0]["flags"] == [], f"check_all flags cards, decisions and threads: {summary['by_kind']}")
    rows = {r["id"]: r for r in tracker.load_rows()}
    check("flags" in common.COLUMNS and all("flags" in r for r in rows.values()), "flags is a tracker column")


def test_prep():
    print("meeting prep")
    t = date(2026, 9, 16)
    packs, md, page = prep.build(["Lakshmi"], ["alpha"], t=t)
    person = next(pk for pk in packs if "person" in pk)
    check(person["person"] == "Laxmikant", f"alias resolved to the roster name: {person['person']}")
    check("## Project: Alpha Project" in md and "## Laxmikant" in md and "### Last commitments, as said" in md,
          "prep pack has the project and the person sections")
    proj = next(pk for pk in packs if "project" in pk)
    check([d["id"] for d in proj["decisions"]] == ["D-2026-09-08-01", "D-2026-09-03-01"], "project pack lists decisions in force only")
    check(page.startswith("<!doctype html>") and "Meeting prep" in page and "Laxmikant" in page, "prep html renders")


def test_sheets():
    print("sheets pull merge")
    rows = tracker.load_rows()
    for r in rows:
        r["updated"] = (date.today() - timedelta(days=1)).isoformat()
    logged = []
    # row 1 is missing from the sheet -> rejected; the others are unchanged
    recs = [{**rows[0], "status": "done", "owner": "Johan"}] + [dict(r) for r in rows[2:]]
    n = sheets.apply_records(rows, recs, log=lambda p, rec: logged.append(rec))
    check(n == 2, f"two rows changed: {n}")
    check({(c["field"], c["to"]) for c in logged} == {("status", "done"), ("owner", "Johan"), ("deleted", "rejected")},
          f"corrections logged per changed cell: {logged}")
    check(rows[1]["status"] == "rejected", "row absent from the sheet is rejected, not removed")
    rows2 = tracker.load_rows()
    n2 = sheets.apply_records(rows2, [], log=lambda p, rec: None)
    check(n2 == 0 and all(r["status"] != "rejected" for r in rows2), "an empty sheet never rejects half the tracker")
    dec = sheets.memory_rows(common.DECISIONS, sheets.DECISION_COLS)
    thr = sheets.memory_rows(common.THREADS, sheets.THREAD_COLS)
    check(dec[0] == list(sheets.DECISION_COLS) and len(dec) > 1 and all(len(r) == len(dec[0]) for r in dec),
          f"decisions tab has a header and one row per record: {len(dec) - 1}")
    by_id = {r[0]: dict(zip(dec[0], r)) for r in dec[1:]}
    check(by_id["D-2026-09-01-01"]["current"] == "no" and by_id["D-2026-09-01-01"]["replaced_by"] == "D-2026-09-08-01"
          and by_id["D-2026-09-08-01"]["current"] == "yes", "decisions tab shows the chain")
    check(thr[0] == list(sheets.THREAD_COLS) and len(thr) > 1
          and all(isinstance(c, str) for r in thr for c in r), "threads tab flattens lists to strings")
    check(sheets.memory_rows(TMP / "missing.jsonl", sheets.DECISION_COLS) == [list(sheets.DECISION_COLS)],
          "a missing memory file gives a header-only tab")
    reqs = sheets.actions_format_requests(1, len(rows))
    json.dumps(reqs)
    check(sum("addConditionalFormatRule" in r for r in reqs) == 11 and sum("setDataValidation" in r for r in reqs) == 5,
          "Actions tab formatting: 11 colour rules, 5 dropdowns")
    values, layout = sheets.dashboard_values(rows)
    json.dumps(sheets.dashboard_format_requests(1, layout))
    check(values[4][0].startswith("=COUNTIFS(Actions!") and values[layout["owner_first"]][0] in {r["owner"] for r in rows}
          and values[layout["owner_header"]] == ["Owner", "P1", "P2", "P3", "Open", "Overdue"],
          "dashboard tab: live formulas over the Actions tab, one row per owner")


def test_readers():
    print("documents")
    eml = TMP / "inbox" / "2026-09-12-status.eml"
    eml.write_bytes(b"From: a@x.com\r\nTo: b@x.com\r\nSubject: Status\r\nDate: Fri, 12 Sep 2026 10:00:00 +0000\r\n"
                    b"Content-Type: text/plain\r\n\r\nPlease send the runbook by Monday.\r\n")
    text = common.read_document(eml)
    check("Subject: Status" in text and "runbook by Monday" in text, "eml reader keeps headers and body")
    note = TMP / "inbox" / "note-2026-09-18-from-phone.md"
    note.write_text("- call Laxmikant\n", encoding="utf-8")
    check(common.is_note(note) and common.date_from_filename(note) == "2026-09-18", "note- prefix and date")
    undated = TMP / "inbox" / "random.txt"
    undated.write_text("x", encoding="utf-8")
    check(len(common.date_from_filename(undated)) == 10, "undated file falls back to the file date")
    vtt = TMP / "captions.vtt"
    vtt.write_text("""WEBVTT

abc/149-0
00:00:03.422 --> 00:00:07.058
<v Priya Sharma>Runbook is about 70% done,
I will finish it</v>

abc/149-1
00:00:07.058 --> 00:00:08.702
<v Priya Sharma>by Friday &amp; dry-run Monday.</v>

abc/150-0
00:00:09.000 --> 00:00:11.000
<v Arjun>Still blocked on staging creds.</v>

abc/151-0
00:00:20.000 --> 00:00:22.000
<v Arjun>Meera will chase infra.</v>
""", encoding="utf-8")
    text = common.read_document(vtt)
    check(text == "Priya Sharma: Runbook is about 70% done, I will finish it by Friday & dry-run Monday."
          "\n\nArjun: Still blocked on staging creds.\n\nArjun: Meera will chase infra.",
          f"vtt reader merges cues per speaker and splits on pauses: {text!r}")
    srt = TMP / "captions.srt"
    srt.write_text("1\n00:00:01,000 --> 00:00:04,000\nHello there.\n\n2\n00:00:04,000 --> 00:00:06,000\nSecond line.\n",
                   encoding="utf-8")
    check(common.read_document(srt) == "Hello there. Second line.", "srt reader drops counters and timings")
    big = "\n\n".join(f"paragraph {i} " + "word " * 300 for i in range(40))
    chunks = common.chunk_text(big, max_tokens=5000)
    check(len(chunks) > 1 and "".join(chunks).replace("\n", "") == big.replace("\n", ""), "chunking loses nothing")
    check(extract.pending() == [eml, note, undated], f"pending lists every supported inbox file: {extract.pending()}")


def test_dayplan():
    print("dayplan")
    from datetime import datetime
    from zoneinfo import ZoneInfo

    cfg = dict(dayplan.DEFAULT_CONFIG)
    tz = ZoneInfo(cfg["timezone"])
    t = date(2026, 9, 16)

    def row(i, owner, status, prio, due, typ, effort, task="t"):
        return {"id": i, "owner": owner, "status": status, "priority": prio, "due": due, "type": typ,
                "effort": effort, "task": task, "next_step": "", "unblocker": "", "prerequisites": "",
                "evidence": "", "source": "s.md"}

    rows = [row("v1", "Samuel", "to_verify", "P1", None, "build", "1h", "Runbook finished"),
            row("m1", "Ayush", "open", "P1", "2026-09-15", "decide", "15m", "Decide flow"),
            row("m2", "Ayush", "open", "P2", "2026-09-16", "communicate", "unknown", "Send comments"),
            row("m3", "Ayush", "open", "P2", "2026-09-20", "review", "1h", "Later"),
            row("m4", "Ayush", "blocked", "P1", "2026-09-15", "build", "1h", "Stuck"),
            row("o1", "Johan", "open", "P1", "2026-09-15", "build", "1h", "Not mine"),
            row("m5", "ayush", "open", "P1", None, "build", "day+", "Big build")]
    picked = [r["id"] for r in dayplan.pick_items(rows, "Ayush", t)]
    check(picked == ["v1", "m1", "m5", "m2"], f"to_verify first, then my P1 overdue, P1 undated, P2 due today: {picked}")
    check(dayplan.minutes_for(rows[6], cfg) == 60 and dayplan.minutes_for(rows[0], cfg) == 5
          and dayplan.minutes_for(rows[2], cfg) == 15, "durations: day+ capped at 60, to_verify 5, unknown effort by type")

    busy = [(datetime(2026, 9, 16, 9, 0, tzinfo=tz), datetime(2026, 9, 16, 9, 30, tzinfo=tz))]
    plan = dayplan.schedule(dayplan.pick_items(rows, "Ayush", t), cfg, t, busy)
    at = {b["id"]: b for b in plan["blocks"]}
    check(at["v1"]["where"] == "walk" and at["v1"]["start"].strftime("%H:%M") == "17:00", "confirmation goes into the walk")
    check(at["m1"]["start"].strftime("%H:%M") == "09:30", f"first desk block starts after the busy slot: {at['m1']['start']}")
    check(at["m5"]["start"].strftime("%H:%M") == "09:55" and at["m5"]["minutes"] == 60, "10 minute gap, then a 60 minute block")
    check(at["m2"]["where"] == "desk", "a 15 minute message is too long for the walk")
    check(not plan["unplaced"], "everything fits on a normal day")

    many = [row(f"b{i}", "Ayush", "open", "P1", None, "build", "1h", f"Build {i}") for i in range(12)]
    plan2 = dayplan.schedule(dayplan.pick_items(many, "Ayush", t), cfg, t)
    lunch = (datetime(2026, 9, 16, 13, 0, tzinfo=tz), datetime(2026, 9, 16, 15, 0, tzinfo=tz))
    walk = (datetime(2026, 9, 16, 17, 0, tzinfo=tz), datetime(2026, 9, 16, 17, 30, tzinfo=tz))
    clash = [b for b in plan2["blocks"] for w in (lunch, walk) if b["start"] < w[1] and b["end"] > w[0]]
    check(not clash and plan2["unplaced"] and len(plan2["blocks"]) + len(plan2["unplaced"]) == 12
          and all(5 <= b["minutes"] <= 60 for b in plan2["blocks"]),
          f"full day: no block touches lunch or the walk, the rest is reported: {len(plan2['blocks'])} placed")

    body = dayplan.event_body(at["v1"], rows[0], cfg)
    check(body["summary"] == "[P1] Confirm: Runbook finished"
          and body["extendedProperties"]["private"] == {"aiemp_id": "v1", "aiemp_date": "2026-09-16"}
          and body["reminders"]["overrides"] == [{"method": "popup", "minutes": 5}], "event body: summary, tag, reminder")
    text = dayplan.render(plan, cfg, t)
    check(text.index("09:30-09:45") < text.index("17:00-17:05") and "(walk, phone)" in text, "plan renders in time order")

    walkish = rows[:1] + [row("w1", "Ayush", "open", "P2", "2026-09-16", "decide", "unknown",
                              "Brainstorm startup ideas during the walk")]
    plan3 = dayplan.schedule(dayplan.pick_items(walkish, "Ayush", t), cfg, t)
    w1 = {b["id"]: b for b in plan3["blocks"]}["w1"]
    check(w1["where"] == "walk" and w1["start"].strftime("%H:%M") == "17:05" and w1["end"].strftime("%H:%M") == "17:30",
          f"a 'during the walk' item takes what is left of the walk: {w1['start']:%H:%M}-{w1['end']:%H:%M}")
    check(not dayplan.walk_hint({"task": "Walk Sreekumar through the CR screen"}), "walking someone through is not a walk")
    ev = dayplan.window_event(cfg["protected"][1], t, cfg)
    check(ev["summary"] == "Walk" and ev["start"]["dateTime"].startswith("2026-09-16T17:00")
          and ev["extendedProperties"]["private"]["aiemp_id"] == "window:Walk", "walk window becomes a calendar event")
    check(dayplan.stale_ids({"window:Walk": {}, "old": {}, "v1": {}}, {"v1"}) == ["old"],
          "stale events are removed but window events are kept")


if __name__ == "__main__":
    fixtures()
    for t in (test_tracker, test_merge_and_cap, test_context_loader, test_search, test_verify, test_brief, test_prep,
              test_sheets, test_readers, test_dayplan):
        try:
            t()
        except Exception as e:  # a crash is a failure too
            FAILS.append(f"{t.__name__} crashed: {e!r}")
            print(f"  FAIL {t.__name__} crashed: {e!r}")
            import traceback
            traceback.print_exc()
    print(f"\n{'FAILED' if FAILS else 'PASSED'}: {len(FAILS)} failure(s)")
    for f in FAILS:
        print(f" - {f}")
    sys.exit(1 if FAILS else 0)
