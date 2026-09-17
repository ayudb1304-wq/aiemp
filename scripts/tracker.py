"""Excel tracker: the single source of truth for action items.

Usage:
  python scripts/tracker.py cards/2026-09-16-intelligent-tracker-sync.json [...]
  python scripts/tracker.py --close <id>          # mark done from the CLI (a human action)

Rules the code enforces:
- the agent never sets `done`; a card that says done becomes `to_verify`
- a due date change is recorded in notes as "due <old> -> <new>" so slips can be counted
- re-running the same cards file creates no duplicates (same source + same task text)
"""
import json
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from common import (COLUMNS, DATE_COLUMNS, EFFORTS, OPEN_STATUSES, ORIGINS, TRACKER, TYPES, UNASSIGNED,
                    parse_date, slug, today)

FONT = "Arial"
NAVY = "1F3864"
INK = "1F2933"
MUTED = "6B7280"
TILE = "F3F6FB"
LINE = "D9DEE7"
# Sequential blue steps for P1 > P2 > P3 (one hue, dark -> light) and the fixed status colours.
PRIORITY_COLOURS = {"P1": "1C5CAB", "P2": "5598E7", "P3": "B7D3F6"}
CHART_BLUE = "2A78D6"
FILL_P1 = PatternFill("solid", fgColor="FBE3E3")
FILL_P2 = PatternFill("solid", fgColor="FFF4D6")
FILL_P3 = PatternFill("solid", fgColor="EEF1F5")
FILL_DONE = PatternFill("solid", fgColor="E2F3E2")
FILL_VERIFY = PatternFill("solid", fgColor="FFF2CC")
FILL_BLOCKED = PatternFill("solid", fgColor="FDE4D7")
FILL_PROGRESS = PatternFill("solid", fgColor="E1ECFB")
FILL_HEADER = PatternFill("solid", fgColor=NAVY)
FILL_TILE = PatternFill("solid", fgColor=TILE)
WIDTHS = {"task": 58, "owner": 14, "priority": 9, "status": 12, "due": 11, "project": 18, "team": 16,
          "type": 12, "next_step": 40, "unblocker": 13, "effort": 9, "blocked_by": 20,
          "prerequisites": 26, "notes": 40, "evidence": 45, "basis": 26, "meeting": 24,
          "origin": 10, "updated": 11, "closed": 11, "created": 11, "source": 32, "id": 32}
MAX_ROWS = 5000   # conditional formats and dropdowns cover this many rows
ADVISOR_FIELDS = ("type", "next_step", "prerequisites", "unblocker", "effort", "basis")
DUE_MARK = re.compile(r"due (\d{4}-\d{2}-\d{2}|none) -> (\d{4}-\d{2}-\d{2}|none)")


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, list):
        return "; ".join(str(x) for x in v if x)
    return str(v)


def load_rows() -> list[dict]:
    if not TRACKER.exists():
        return []
    wb = load_workbook(TRACKER)
    ws = wb["Actions"] if "Actions" in wb.sheetnames else wb.active
    headers = [c.value for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(values):
            continue
        row = dict(zip(headers, values))
        for k in ("created", "due", "updated", "closed"):
            d = parse_date(row.get(k))
            row[k] = d.isoformat() if d else ""
        rows.append({c: _s(row.get(c)) for c in COLUMNS})
    return rows


def _sorted(rows: list[dict]) -> list[dict]:
    """Open first, then P1 -> P3, then earliest due. Shared with the Sheet push."""
    order = {"P1": 0, "P2": 1, "P3": 2}
    return sorted(rows, key=lambda r: (r["status"] not in OPEN_STATUSES,
                                       order.get(r["priority"], 9), r["due"] or "9999"))


def _col(name: str) -> str:
    return get_column_letter(COLUMNS.index(name) + 1)


def _cell_value(col: str, v):
    """Dates are written as real dates (sortable, filterable by month in Excel); the rest as text."""
    if col in DATE_COLUMNS:
        return parse_date(v)
    return _s(v)


def _write_actions(ws, rows: list[dict]) -> None:
    thin = Side(style="thin", color=LINE)
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[1].height = 24
    for r in rows:
        ws.append([_cell_value(c, r.get(c, "")) for c in COLUMNS])
        for c in ws[ws.max_row]:
            c.font = Font(name=FONT, size=10, color=INK)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=thin)
    for col in DATE_COLUMNS:
        for c in ws[_col(col)][1:]:
            c.number_format = "yyyy-mm-dd"
    for col in ("priority", "status", "due", "effort", "origin", "updated", "closed", "created"):
        for c in ws[_col(col)][1:]:
            c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for i, col in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(col, 14)
    ws.freeze_panes = "B2"          # header row and the task column stay in view
    ws.sheet_view.zoomScale = 90

    last = get_column_letter(len(COLUMNS))
    n = max(len(rows), 1)
    table = Table(displayName="Actions", ref=f"A1:{last}{n + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True,
                                          showFirstColumn=False, showLastColumn=False, showColumnStripes=False)
    ws.add_table(table)

    # Colour is live: change a status or priority in Excel and the row restyles itself.
    rng = f"A2:{last}{MAX_ROWS}"
    st, pr, du = (f"${_col(c)}2" for c in ("status", "priority", "due"))
    closed = f'OR({st}="done",{st}="rejected")'
    ws.conditional_formatting.add(f"{_col('status')}2:{_col('status')}{MAX_ROWS}", FormulaRule(
        formula=[f'{st}="done"'], fill=FILL_DONE, font=Font(name=FONT, size=10, bold=True, color="1E7B1E")))
    for value, fill, colour in (("to_verify", FILL_VERIFY, "8A6100"), ("blocked", FILL_BLOCKED, "9A3F12"),
                                ("in_progress", FILL_PROGRESS, "1C5CAB")):
        ws.conditional_formatting.add(f"{_col('status')}2:{_col('status')}{MAX_ROWS}", FormulaRule(
            formula=[f'{st}="{value}"'], fill=fill, font=Font(name=FONT, size=10, bold=True, color=colour)))
    for value, fill, colour in (("P1", FILL_P1, "9B1C1C"), ("P2", FILL_P2, "8A6100"), ("P3", FILL_P3, MUTED)):
        ws.conditional_formatting.add(f"{_col('priority')}2:{_col('priority')}{MAX_ROWS}", FormulaRule(
            formula=[f'AND({pr}="{value}",NOT({closed}))'], fill=fill,
            font=Font(name=FONT, size=10, bold=True, color=colour)))
    due_rng = f"{_col('due')}2:{_col('due')}{MAX_ROWS}"
    ws.conditional_formatting.add(due_rng, FormulaRule(
        formula=[f'AND({du}<>"",{du}<TODAY(),NOT({closed}))'], font=Font(name=FONT, size=10, bold=True, color="C0262A")))
    ws.conditional_formatting.add(due_rng, FormulaRule(
        formula=[f'AND({du}<>"",{du}=TODAY(),NOT({closed}))'], font=Font(name=FONT, size=10, bold=True, color="8A6100")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[closed], font=Font(name=FONT, size=10, color=MUTED)))

    # Dropdowns keep hand edits inside the vocabulary the scripts understand.
    for col, values in (("status", sorted(OPEN_STATUSES) + ["done", "rejected"]), ("priority", ["P1", "P2", "P3"]),
                        ("type", TYPES), ("effort", EFFORTS), ("origin", ORIGINS)):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        dv.error, dv.errorTitle = f"Pick one of: {', '.join(values)}", f"Unknown {col}"
        ws.add_data_validation(dv)
        dv.add(f"{_col(col)}2:{_col(col)}{MAX_ROWS}")


def _horizon(r: dict, t: date) -> str:
    d = parse_date(r["due"])
    if not d:
        return "No date"
    if d < t:
        return "Overdue"
    if d == t:
        return "Due today"
    if (d - t).days <= 7:
        return "Next 7 days"
    return "Later"


def _write_dashboard(ws, rows: list[dict], t: date) -> None:
    """Summary of the Actions sheet as live formulas (COUNTIFS over the Actions columns), so a
    status or priority changed in Excel moves the tiles and charts immediately. The owner and
    project lists are written from the current rows and refresh on the next run. Excel computes
    the formulas on open (fullCalcOnLoad); GitHub's file preview shows them blank."""
    open_rows = [r for r in rows if r["status"] in OPEN_STATUSES]
    thin = Side(style="thin", color=LINE)
    S, P, D, O, J = (f"Actions!${_col(c)}$2:${_col(c)}${MAX_ROWS}" for c in ("status", "priority", "due", "owner", "project"))
    is_open = S + ',{"' + '","'.join(sorted(OPEN_STATUSES)) + '"}'   # array constant: one COUNTIFS per status

    def count(*criteria) -> str:
        return "=SUMPRODUCT(COUNTIFS(" + ",".join((is_open,) + criteria) + "))"

    def put(ref, value, *, bold=False, size=10, colour=INK, fill=None, align="left", fmt=None):
        c = ws[ref]
        c.value = value
        c.font = Font(name=FONT, bold=bold, size=size, color=colour)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt
        return c

    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 24, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11, "I": 3}
    widths.update({get_column_letter(i): 11 for i in range(10, 18)})
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    put("B1", "Action tracker", bold=True, size=20, colour=NAVY)
    put("B2", f"Live: every number is a formula over the Actions sheet, so a status changed there moves this "
              f"page at once. Owner and project lists refresh on each run (last {t.isoformat()}). Edits that "
              f"should stick belong in the Google Sheet.", colour=MUTED, size=9)
    ws.merge_cells("B2:Q2")

    tiles = [("OPEN ITEMS", count()), ("P1 OPEN", count(P, '"P1"')),
             ("OVERDUE", count(D, '"<"&TODAY()')), ("DUE TODAY", count(D, "TODAY()")),
             ("TO VERIFY", f'=COUNTIF({S},"to_verify")'), ("BLOCKED", f'=COUNTIF({S},"blocked")'),
             ("DONE", f'=COUNTIF({S},"done")')]
    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 34
    for col, (label, value) in zip("BCDEFGH", tiles):
        put(f"{col}4", label, size=8, bold=True, colour=MUTED, fill=FILL_TILE, align="center")
        put(f"{col}5", value, size=22, bold=True, colour=NAVY, fill=FILL_TILE, align="center")
    for col in ("C", "D"):   # P1 open and overdue turn red when non-zero
        ws.conditional_formatting.add(f"{col}5", FormulaRule(formula=[f"{col}5>0"],
                                                             font=Font(name=FONT, size=22, bold=True, color="C0262A")))
        ws[f"{col}4"].border = Border(left=thin, right=thin, top=thin)
        ws[f"{col}5"].border = Border(left=thin, right=thin, bottom=thin)

    def header(row, cells):
        for ref, text in cells:
            put(ref, text, bold=True, size=9, colour="FFFFFF", fill=FILL_HEADER,
                align="left" if ref[0] == "B" else "center")
        ws.row_dimensions[row].height = 18

    def section(row, title, note=""):
        put(f"B{row}", title, bold=True, size=12, colour=NAVY)
        if note:
            put(f"D{row}", note, size=9, colour=MUTED)
            ws.merge_cells(f"D{row}:Q{row}")

    # Open items by owner (feeds the workload chart)
    r0 = 8
    section(r0, "Open items by owner", "stacked by priority; the chart on the right reads from this table")
    header(r0 + 1, [(f"B{r0 + 1}", "Owner"), (f"C{r0 + 1}", "P1"), (f"D{r0 + 1}", "P2"), (f"E{r0 + 1}", "P3"),
                    (f"F{r0 + 1}", "Open"), (f"G{r0 + 1}", "Overdue")])
    owners = Counter(r["owner"] or "Unassigned" for r in open_rows)
    row = r0 + 2
    first_owner_row = row
    for owner, n in owners.most_common():
        who = f"{O},$B{row}"
        vals = [owner, count(who, P, '"P1"'), count(who, P, '"P2"'), count(who, P, '"P3"'), count(who),
                count(who, D, '"<"&TODAY()')]
        for col, v in zip("BCDEFG", vals):
            put(f"{col}{row}", v, align="left" if col == "B" else "center", bold=(col == "F"))
            ws[f"{col}{row}"].border = Border(bottom=thin)
        row += 1
    last_owner_row = max(row - 1, first_owner_row)
    if not owners:
        put(f"B{row}", "No open items", colour=MUTED)
        row += 1
    put(f"B{row}", "Total", bold=True)
    for col in "CDEFG":
        put(f"{col}{row}", f"=SUM({col}{first_owner_row}:{col}{last_owner_row})", bold=True, align="center")
        ws[f"{col}{row}"].border = Border(top=Side(style="medium", color=NAVY))
    ws[f"B{row}"].border = Border(top=Side(style="medium", color=NAVY))
    total_row = row

    if owners:
        chart = BarChart()
        chart.type = "bar"
        chart.grouping = "stacked"
        chart.overlap = 100
        chart.gapWidth = 60
        chart.title = "Open items by owner and priority"
        chart.style = 10
        data = Reference(ws, min_col=3, max_col=5, min_row=r0 + 1, max_row=last_owner_row)
        cats = Reference(ws, min_col=2, min_row=first_owner_row, max_row=last_owner_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        for ser, key in zip(chart.series, ("P1", "P2", "P3")):
            ser.graphicalProperties.solidFill = PRIORITY_COLOURS[key]
            ser.graphicalProperties.line.solidFill = "FFFFFF"
        chart.y_axis.majorGridlines = None
        chart.y_axis.delete = False
        chart.x_axis.delete = False
        chart.x_axis.scaling.orientation = "maxMin"   # busiest owner on top, same order as the table
        chart.legend.position = "b"
        chart.height = 0.55 * max(len(owners), 4) + 3.2
        chart.width = 16
        ws.add_chart(chart, f"J{r0}")

    # Due horizon (feeds the second chart)
    r1 = max(total_row + 3, r0 + 14)
    section(r1, "When is it due", "open items only")
    header(r1 + 1, [(f"B{r1 + 1}", "Horizon"), (f"C{r1 + 1}", "Items"), (f"D{r1 + 1}", "P1")])
    horizons = [("Overdue", (D, '"<"&TODAY()')), ("Due today", (D, "TODAY()")),
                ("Next 7 days", (D, '">"&TODAY()', D, '"<="&TODAY()+7')), ("Later", (D, '">"&TODAY()+7')),
                ("No date", (D, '""'))]
    for i, (h, crit) in enumerate(horizons):
        put(f"B{r1 + 2 + i}", h)
        put(f"C{r1 + 2 + i}", count(*crit), align="center", bold=True)
        put(f"D{r1 + 2 + i}", count(*crit, P, '"P1"'), align="center")
        for col in "BCD":
            ws[f"{col}{r1 + 2 + i}"].border = Border(bottom=thin)
    hz = BarChart()
    hz.type = "col"
    hz.title = "Open items by due horizon"
    hz.style = 10
    hz.add_data(Reference(ws, min_col=3, min_row=r1 + 2, max_row=r1 + 6), titles_from_data=False)
    hz.set_categories(Reference(ws, min_col=2, min_row=r1 + 2, max_row=r1 + 6))
    hz.series[0].graphicalProperties.solidFill = CHART_BLUE
    hz.series[0].graphicalProperties.line.noFill = True
    hz.dataLabels = DataLabelList()
    hz.dataLabels.showVal = True
    hz.legend = None
    hz.y_axis.majorGridlines = None
    hz.y_axis.delete = True
    hz.x_axis.delete = False
    hz.gapWidth = 80
    hz.height, hz.width = 7, 16
    ws.add_chart(hz, f"J{r1}")

    # By project and by status: tables only
    r2 = r1 + 9
    section(r2, "By project")
    header(r2 + 1, [(f"B{r2 + 1}", "Project"), (f"C{r2 + 1}", "Open"), (f"D{r2 + 1}", "P1"),
                    (f"E{r2 + 1}", "Overdue"), (f"F{r2 + 1}", "Done")])
    projects = sorted({r["project"] or UNASSIGNED for r in rows})
    for i, p in enumerate(projects):
        here = f"{J},$B{r2 + 2 + i}"
        vals = [p, count(here), count(here, P, '"P1"'), count(here, D, '"<"&TODAY()'),
                f'=COUNTIFS({here},{S},"done")']
        for col, v in zip("BCDEF", vals):
            put(f"{col}{r2 + 2 + i}", v, align="left" if col == "B" else "center", bold=(col == "C"))
            ws[f"{col}{r2 + 2 + i}"].border = Border(bottom=thin)
    r3 = r2 + 2 + max(len(projects), 1) + 2
    section(r3, "By status")
    header(r3 + 1, [(f"B{r3 + 1}", "Status"), (f"C{r3 + 1}", "Items")])
    statuses = sorted(OPEN_STATUSES) + ["done", "rejected"]
    for i, st in enumerate(statuses):
        put(f"B{r3 + 2 + i}", st)
        put(f"C{r3 + 2 + i}", f'=COUNTIF({S},$B{r3 + 2 + i})', align="center", bold=True)
        for col in "BC":
            ws[f"{col}{r3 + 2 + i}"].border = Border(bottom=thin)
    r4 = r3 + 2 + len(statuses) + 1
    put(f"B{r4}", "Reading the Actions sheet: rows are ordered open first, then P1 to P3, then by due date. "
                  "Red priority = P1; amber status = reported complete, waiting for your confirmation "
                  "(to_verify); orange = blocked; green = done; a red due date is overdue. Only you set done.",
        colour=MUTED, size=9)
    ws.merge_cells(f"B{r4}:Q{r4 + 1}")
    ws.row_dimensions[r4].height = 16
    ws.row_dimensions[r4 + 1].height = 16


def save_rows(rows: list[dict]) -> None:
    t = date.today()
    for r in rows:
        if r.get("status") == "done" and not r.get("closed"):
            r["closed"] = today()
    rows = _sorted(rows)
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    actions = wb.create_sheet("Actions")
    _write_actions(actions, rows)
    _write_dashboard(dash, rows, t)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True   # the dashboard formulas have no cached values
    TRACKER.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TRACKER)


def open_items(rows: list[dict] | None = None, project: str | None = None) -> list[dict]:
    rows = load_rows() if rows is None else rows
    out = [r for r in rows if r["status"] in OPEN_STATUSES]
    if project:
        out = [r for r in out if slug(r.get("project")) == slug(project)]
    return out


def _next_id(rows: list[dict], meeting_date: str, project: str) -> str:
    prefix = f"{meeting_date}-{slug(project or UNASSIGNED)}-"
    n = sum(1 for r in rows if r["id"].startswith(prefix)) + 1
    return f"{prefix}{n:02d}"


def _add_note(r: dict, meeting_date: str, note: str) -> None:
    r["notes"] = (r["notes"] + " | " if r["notes"] else "") + f"{meeting_date}: {note}"


def slip_count(r: dict) -> tuple[int, str]:
    """(number of due changes, first due) parsed from the notes markers."""
    marks = DUE_MARK.findall(r.get("notes") or "")
    first = next((m[0] for m in marks if m[0] != "none"), r.get("due") or "")
    return len(marks), first


def upsert(cards: list[dict], meeting: str, meeting_date: str, source: str,
           project: str | None = None, origin: str = "planned") -> tuple[list[str], list[str]]:
    """Merge extracted cards into the tracker. Returns (created_ids, updated_ids)."""
    rows = load_rows()
    by_id = {r["id"]: r for r in rows}
    created, updated = [], []
    for c in cards:
        match = c.get("matches_existing_id")
        status = c.get("status") or "open"
        if status == "done":  # the agent never closes an item
            status = "to_verify"
        if match and match in by_id:
            r = by_id[match]
            note = c.get("update_note") or f"Mentioned again in {meeting}"
            if status == "to_verify" and c.get("evidence"):
                note = f"{note} [{c['evidence']}]"
            if f"{meeting_date}: {note}" in r["notes"]:
                continue  # same file processed again: this update is already applied
            if c.get("due") and c["due"] != r["due"]:
                _add_note(r, meeting_date, f"due {r['due'] or 'none'} -> {c['due']}")
                r["due"] = c["due"]
            for k in ("priority", "owner", "blocked_by"):
                if c.get(k):
                    r[k] = _s(c[k])
            if c.get("status"):
                r["status"] = status
            if c.get("basis") and _s(c["basis"]) != r.get("basis"):
                for k in ADVISOR_FIELDS:  # refresh the advice when its basis changed
                    if c.get(k) is not None:
                        r[k] = _s(c[k])
            elif c.get("type") and not r.get("type"):
                r["type"] = _s(c["type"])
            _add_note(r, meeting_date, note)
            r["updated"] = today()
            updated.append(r["id"])
            continue
        dup = next((r for r in rows if r["source"] == source and r["task"] == c.get("task", "")), None)
        if dup:  # same file processed again: idempotent
            continue
        proj = c.get("project") or project or UNASSIGNED
        row = {col: "" for col in COLUMNS}
        row.update({
            "id": _next_id(rows, meeting_date, proj),
            "created": meeting_date,
            "meeting": meeting,
            "team": c.get("team", ""),
            "project": slug(proj),
            "owner": c.get("owner", ""),
            "task": c.get("task", ""),
            "due": c.get("due") or "",
            "priority": c.get("priority", "P2"),
            "status": status,
            "blocked_by": _s(c.get("blocked_by")),
            "evidence": c.get("evidence", ""),
            "source": source,
            "updated": today(),
            "origin": c.get("origin") or origin,
        })
        for k in ADVISOR_FIELDS:
            row[k] = _s(c.get(k))
        if status == "to_verify" and c.get("evidence"):
            _add_note(row, meeting_date, f"Reported complete [{c['evidence']}]")
        rows.append(row)
        by_id[row["id"]] = row
        created.append(row["id"])
    save_rows(rows)
    return created, updated


def close(item_id: str) -> None:
    rows = load_rows()
    for r in rows:
        if r["id"] == item_id:
            r["status"], r["updated"], r["closed"] = "done", today(), today()
    save_rows(rows)


if __name__ == "__main__":
    args = sys.argv[1:]
    if args[:1] == ["--close"]:
        close(args[1])
        print(f"closed {args[1]}")
        sys.exit()
    for p in args:
        data = json.loads(Path(p).read_text(encoding="utf-8"))
        c, u = upsert(data["cards"], data["meeting"], data["date"], data["source"],
                      project=data.get("project"), origin=data.get("origin", "planned"))
        print(f"{p}: {len(c)} created, {len(u)} updated")
